import openpyxl
import pytest

from services.deficit_payment_table import (
    attach_deficit_payment_plan,
    build_deficit_payment_by_department,
    build_deficit_payment_fulfillment_board,
    build_deficit_payment_receipt_table,
    build_deficit_payment_share,
    parse_payment_receipt_plan_sheet,
)
from services.dashboard_engine import render_spec
from models.dashboard_spec import DashboardSpec, Tab, Tile, TileSource
from services.report_profiles.deficit_profile import deficit_payment_tab_tiles
from tests.test_deficit_sales_table import _sample_df


def test_parse_plan_from_third_sheet(tmp_path):
    path = tmp_path / "deficit.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Дефицит"
    ws1["A1"] = "заголовок"
    wb.create_sheet("Прогноз")
    ws3 = wb.create_sheet("План")
    ws3["B3"] = "Службы"
    ws3["C3"] = "План поступлений, р."
    ws3["B4"] = "СТО"
    ws3["C4"] = 100.0
    ws3["B5"] = "Совтест АТЕ"
    ws3["C5"] = 250.0
    wb.save(path)

    parsed, company = parse_payment_receipt_plan_sheet(str(path))
    assert parsed["СТО"] == pytest.approx(100.0)
    assert company == pytest.approx(250.0)


def _payment_df_with_plan(tmp_path):
    path = tmp_path / "deficit.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Дефицит"
    wb.create_sheet("Прогноз")
    ws3 = wb.create_sheet("План")
    ws3["B2"] = "Службы"
    ws3["C2"] = "План поступлений, р."
    ws3["B3"] = "СИО"
    ws3["C3"] = 50.0
    ws3["B4"] = "ОАПЛиС"
    ws3["C4"] = 30.0
    ws3["B5"] = "СТО"
    ws3["C5"] = 120.0
    ws3["B6"] = "Совтест АТЕ"
    ws3["C6"] = 200.0
    wb.save(path)

    df = _sample_df()
    df["оплачено по заказу"] = [100.0, 200.0, 200.0, 100.0]
    attach_deficit_payment_plan(df, str(path))
    return df


def test_payment_receipt_table_plan_and_fact(tmp_path):
    df = _payment_df_with_plan(tmp_path)
    tile = Tile(
        title="Поступление оплаты",
        chart_type="table",
        source=TileSource(kind="deficit_payment_receipt"),
        agg="sum",
        top_n=50,
        sort="none",
    )
    data = build_deficit_payment_receipt_table(df, tile)
    assert data["plan"] == pytest.approx(200.0)
    assert data["fact"] == pytest.approx(600.0)
    assert data["fulfill_pct"] == 300
    assert data["table"]["columns"] == ["План", "Факт", "Выполнение плана"]
    assert "department_table" not in data

    rendered = render_spec(
        df, DashboardSpec(tabs=[Tab(title="Структура", tiles=[tile])])
    )["tabs"][0]["tiles"][0]
    assert rendered["table"]["layout"] == "payment_receipt"
    assert rendered["table"]["rows"][0]["values"][0] == pytest.approx(200.0)
    assert rendered.get("plotly_json")
    import json

    fig = json.loads(rendered["plotly_json"])
    assert fig["layout"]["title"]["text"] == "ПЛАН И ФАКТ ОПЛАТЫ, РУБ."


def test_payment_tiles_split_like_sales(tmp_path):
    df = _payment_df_with_plan(tmp_path)
    tiles = deficit_payment_tab_tiles("quarter")
    assert [t.source.kind for t in tiles] == [
        "deficit_payment_receipt",
        "deficit_payment_by_department",
        "deficit_payment_fulfillment_board",
        "deficit_payment_share",
    ]
    rendered = render_spec(df, DashboardSpec(tabs=[Tab(title="Структура", tiles=tiles)]))
    out = rendered["tabs"][0]["tiles"]
    assert len(out) == 4
    assert out[0]["table"]["columns"][0] == "План"
    assert out[0].get("plotly_json")

    dept = out[1]["table"]
    assert dept["index_label"] == "Службы"
    assert dept["columns"][0] == "План поступлений, р."
    assert dept["totals_label"] == "Совтест АТЕ"
    by_label = {row["label"]: row["values"] for row in dept["rows"]}
    assert "СИО" in by_label and "ОАПЛиС" in by_label
    assert "СИО и АПЛиС" not in by_label
    assert out[1].get("plotly_json")

    fulfill = out[2]["table"]
    assert fulfill["columns"] == ["Оплата, р."]
    assert out[2].get("plotly_json")

    share = out[3]["table"]
    assert share["columns"] == ["Оплата, р.", "Доля поступлений"]
    assert out[3].get("plotly_json")


def test_payment_month_period_splits_all_tables(tmp_path):
    df = _payment_df_with_plan(tmp_path)
    df["заказ клиента"] = [
        "Заказ клиента ЗК-1 от 10.07.2026",
        "Заказ клиента ЗК-2 от 15.08.2026",
        "Заказ клиента ЗК-3 от 20.09.2026",
        "Заказ клиента ЗК-4 от 25.09.2026",
    ]
    tile = Tile(
        title="Поступление оплаты",
        chart_type="table",
        source=TileSource(kind="deficit_payment_receipt", period="month"),
        agg="sum",
        top_n=50,
        sort="none",
    )
    data = build_deficit_payment_receipt_table(df, tile)
    assert len(data["period_bars"]) >= 2
    assert sum(int(b["pct"]) for b in data["period_bars"]) == data["fulfill_pct"]
    assert len(data["table"]["columns"]) >= 4  # план + периоды + %

    tiles = deficit_payment_tab_tiles("month")
    rendered = render_spec(df, DashboardSpec(tabs=[Tab(title="Структура", tiles=tiles)]))
    out = rendered["tabs"][0]["tiles"]
    assert len(out[0]["table"]["columns"]) >= 4
    dept_cols = out[1]["table"]["columns"]
    assert any("июл" in c.lower() or "сен" in c.lower() for c in dept_cols)
    share_cols = out[3]["table"]["columns"]
    assert "Доля поступлений" in share_cols
    assert len(share_cols) >= 3
    # График доли: все месяцы в stack компании, не только первый столбец.
    import json

    from services.dashboard_engine import _render_deficit_payment_share_chart

    share_table = out[3]["table"]
    fig = json.loads(_render_deficit_payment_share_chart(share_table).to_json())
    company = share_table.get("totals_label") or "Совтест АТЕ"
    stacked = [
        t
        for t in fig["data"]
        if t.get("x") == [company] or (isinstance(t.get("x"), list) and company in t.get("x", []))
    ]
    assert len(stacked) >= 2
    period_names = " ".join(str(t.get("name") or "") for t in stacked).lower()
    assert "июл" in period_names or "сен" in period_names or "авг" in period_names
    # Итог компании = сумма периодов, а не колонка «июл».
    assert share_table["order_sum_column_index"] == len(share_cols) - 2


def test_payment_builders_standalone(tmp_path):
    df = _payment_df_with_plan(tmp_path)
    tile = Tile(
        title="t",
        chart_type="table",
        source=TileSource(kind="deficit_payment_by_department", period="quarter"),
        agg="sum",
        top_n=50,
        sort="none",
    )
    assert "table" in build_deficit_payment_by_department(df, tile)
    assert "table" in build_deficit_payment_fulfillment_board(df, tile)
    assert "table" in build_deficit_payment_share(df, tile)
