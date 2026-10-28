"""Именованные отчёты-канвасы в SQLite + пересчёт блоков и Excel."""
from __future__ import annotations

import io
import json
import logging
import uuid
from datetime import datetime, timezone
from typing import Any

import pandas as pd

from models.dashboard_spec import DashboardSpec, Tab, Tile, TileSource
from services import db_service

logger = logging.getLogger(__name__)

MAX_BLOCKS = 24
PERIOD_KINDS = {
    "halfyear",
    "pivot",
    "outcome",
    "deal_statuses",
    "in_work_stages",
    "deals_dynamics",
    "deals_dynamics_departments",
    "deals_dynamics_outcome_share",
    "deals_conversion",
    "deals_money",
    "deficit_sales_by_department",
    "deficit_sales_summary",
    "deficit_sales_forecast_vitrine",
    "deficit_sales_fulfillment_board",
    "deficit_sales_share",
    "deficit_payment_receipt",
    "deficit_payment_by_department",
    "deficit_payment_fulfillment_board",
    "deficit_payment_share",
}


def _now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def list_reports(workspace_id: str | None = None, file_id: str | None = None) -> list[dict]:
    return db_service.list_saved_reports(workspace_id=workspace_id, file_id=file_id)


def get_report(report_id: str) -> dict | None:
    return db_service.get_saved_report(report_id)


def save_report(
    *,
    title: str,
    payload: dict,
    report_id: str | None = None,
    workspace_id: str | None = None,
    file_id: str | None = None,
) -> dict:
    blocks = list(payload.get("blocks") or [])
    if len(blocks) > MAX_BLOCKS:
        raise ValueError(f"В отчёте не больше {MAX_BLOCKS} блоков")
    rid = report_id or uuid.uuid4().hex
    now = _now_iso()
    existing = db_service.get_saved_report(rid) if report_id else None
    created = existing["created_at"] if existing else now
    row = {
        "id": rid,
        "workspace_id": workspace_id,
        "file_id": file_id,
        "title": (title or "Отчёт").strip() or "Отчёт",
        "payload": {"title": title, "blocks": blocks},
        "created_at": created,
        "updated_at": now,
    }
    db_service.upsert_saved_report(row)
    return db_service.get_saved_report(rid) or row


def delete_report(report_id: str) -> bool:
    return db_service.delete_saved_report(report_id)


def _tile_from_source(title: str, source: dict, period: str | None) -> Tile:
    src = dict(source or {})
    kind = str(src.get("kind") or "group")
    if period and kind in PERIOD_KINDS:
        src["period"] = period
    elif period and "period" in src:
        src["period"] = period
    # Убираем None-значения — Pydantic иначе спотыкается на лишних ключах.
    clean = {k: v for k, v in src.items() if v is not None}
    return Tile(
        title=title or "Блок",
        chart_type="table",
        source=TileSource.model_validate(clean),
        agg="sum",
        top_n=50,
        sort="none",
    )


def render_block(
    df: pd.DataFrame,
    *,
    title: str,
    source: dict,
    period: str | None = None,
) -> dict[str, Any]:
    """Пересчитать один блок отчёта (таблица/график) с заданным периодом."""
    from services.dashboard_engine import render_spec

    try:
        tile = _tile_from_source(title, source, period)
    except Exception as exc:
        return {"error": f"Не удалось разобрать источник блока: {exc}"}
    rendered = render_spec(df, DashboardSpec(tabs=[Tab(title="_", tiles=[tile])]))
    out = (rendered.get("tabs") or [{}])[0].get("tiles") or [{}]
    payload = out[0] if out else {}
    if payload.get("error"):
        return {"error": payload["error"]}
    return {
        "title": payload.get("title") or title,
        "table": payload.get("table"),
        "plotly_json": payload.get("plotly_json"),
        "bucket_period": payload.get("bucket_period") or period,
        "chart_type": payload.get("chart_type"),
    }


def export_report_excel(payload: dict, filename: str | None = None) -> bytes:
    """Один лист: блоки сверху вниз по y,x."""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    blocks = sorted(
        list(payload.get("blocks") or []),
        key=lambda b: (
            int((b.get("layout") or {}).get("y") or 0),
            int((b.get("layout") or {}).get("x") or 0),
        ),
    )
    row_i = 1
    title = str(payload.get("title") or filename or "Отчёт")
    ws.cell(row_i, 1, title).font = Font(bold=True, size=14)
    row_i += 2

    for block in blocks:
        kind = block.get("kind") or "chart"
        block_title = str(block.get("title") or "")
        ws.cell(row_i, 1, block_title).font = Font(bold=True, size=12)
        row_i += 1

        if kind == "text":
            text = str(block.get("text") or "")
            for line in text.splitlines() or [""]:
                ws.cell(row_i, 1, line)
                row_i += 1
            row_i += 1
            continue

        table = block.get("table")
        if isinstance(table, dict) and (table.get("columns") or table.get("rows")):
            cols = list(table.get("columns") or [])
            index_label = table.get("index_label") or ""
            headers = ([index_label] if index_label else [""]) + cols
            for c, h in enumerate(headers, start=1):
                ws.cell(row_i, c, h).font = Font(bold=True)
            row_i += 1
            for r in table.get("rows") or []:
                values = [r.get("label"), *(r.get("values") or [])]
                for c, v in enumerate(values, start=1):
                    ws.cell(row_i, c, v)
                row_i += 1
            totals = table.get("totals")
            if totals:
                values = [table.get("totals_label") or "Итого", *totals]
                for c, v in enumerate(values, start=1):
                    cell = ws.cell(row_i, c, v)
                    cell.font = Font(bold=True)
                row_i += 1
            row_i += 1

        plotly_json = block.get("plotly_json")
        if plotly_json and not table:
            png = _plotly_to_png(plotly_json)
            if png:
                img = XLImage(io.BytesIO(png))
                img.width = min(img.width, 640)
                img.height = min(img.height, 360)
                ws.add_image(img, f"A{row_i}")
                row_i += 20
            else:
                ws.cell(row_i, 1, f"[график: {block_title}]")
                row_i += 2
        elif plotly_json and table:
            # таблица уже выведена; график пропускаем на одном листе
            pass

        row_i += 1

    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _plotly_to_png(plotly_json: str) -> bytes | None:
    try:
        import plotly.io as pio
        from plotly.io import from_json

        fig = from_json(plotly_json)
        return pio.to_image(fig, format="png", width=800, height=450, scale=1)
    except Exception as exc:
        logger.info("PNG графика недоступен: %s", exc)
        return None
